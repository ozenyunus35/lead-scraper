"""Write results as CSV, a formatted Excel workbook, and JSON for the web viewer."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")


def write_csv(rows: list[dict], fields: list[str], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:  # BOM: Excel opens UTF-8 correctly
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows: list[dict], fields: list[str], path: Path, sheet: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append([f.replace("_", " ").title() for f in fields])
    for r in rows:
        ws.append([r.get(f, "") for f in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, f in enumerate(fields, 1):
        longest = max([len(str(r.get(f, ""))) for r in rows[:500]] + [len(f)])
        ws.column_dimensions[get_column_letter(i)].width = min(50, max(8, longest + 2))
        if f in ("website", "url", "source_url"):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                c = row[0]
                if c.value:
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
    wb.save(path)


def write_json(rows: list[dict], fields: list[str], meta: dict, path: Path) -> None:
    path.write_text(json.dumps({"meta": meta, "fields": fields, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
